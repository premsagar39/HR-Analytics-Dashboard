"""Prepare HR records, run SQL analysis, and reconcile the reported metrics."""
from pathlib import Path
import argparse,csv,json,sqlite3
from decimal import Decimal
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
FIELDS={'Employee Number':'employee_id','Attrition':'attrition','Department':'department',
        'Job Role':'job_role','Gender':'gender','Age':'age','Education':'education',
        'Education Field':'education_field','Over Time':'overtime','Years At Company':'years_at_company',
        'Job Satisfaction':'job_satisfaction','Monthly Income':'monthly_income','Employee Count':'employee_count'}
NUMERIC={'employee_id','age','years_at_company','job_satisfaction','monthly_income','employee_count'}

def prepare(source):
    book=load_workbook(source,read_only=True,data_only=True)
    sheet=book['HR data'];it=sheet.iter_rows(values_only=True);headers=next(it)
    missing=set(FIELDS)-set(headers)
    if missing:raise ValueError(f'Missing source columns: {sorted(missing)}')
    records=[];seen=set();trimmed=0
    for line,row in enumerate(it,2):
        raw=dict(zip(headers,row));r={}
        for source_name,target in FIELDS.items():
            v=raw[source_name]
            if v is None:raise ValueError(f'Blank {source_name}, source row {line}')
            if target in NUMERIC:
                number=Decimal(str(v))
                if number!=number.to_integral_value():raise ValueError(f'Expected integer {source_name}, row {line}')
                v=int(number)
            else:
                before=str(v);v=before.strip();trimmed+=before!=v
            r[target]=v
        for field in ('attrition','overtime'):
            if r[field].casefold() not in ('yes','no'):raise ValueError(f'Unexpected {field}, row {line}')
            r[field]=r[field].capitalize()
        if r['employee_id'] in seen:raise ValueError(f'Duplicate employee identifier, row {line}')
        seen.add(r['employee_id'])
        if r['employee_count']!=1:raise ValueError('Employee Count must equal one for this row-grain model')
        if not 0<=r['age']<=120 or r['years_at_company']<0:raise ValueError(f'Invalid age or tenure, row {line}')
        if not 1<=r['job_satisfaction']<=4:raise ValueError(f'Invalid satisfaction score, row {line}')
        r['attrition_flag']=int(r['attrition']=='Yes')
        a=r['age'];r['age_band']='Under 25' if a<25 else '25-34' if a<35 else '35-44' if a<45 else '45-54' if a<55 else '55+'
        y=r['years_at_company'];r['tenure_band']='0-1 years' if y<2 else '2-5 years' if y<=5 else '6-10 years' if y<=10 else '11+ years'
        records.append(r)
    book.close()
    return records,{'source_rows':len(records),'retained_rows':len(records),'duplicate_employee_ids':0,'missing_required_values':0,'trimmed_text_values':trimmed}

def load_sqlite(records,path):
    if path.exists():path.unlink()
    con=sqlite3.connect(path)
    cols=list(records[0])
    types={c:'INTEGER' if c in NUMERIC or c=='attrition_flag' else 'TEXT' for c in cols}
    con.execute('CREATE TABLE hr_employees ('+', '.join(f'"{c}" {types[c]} NOT NULL' for c in cols)+', PRIMARY KEY(employee_id))')
    con.executemany('INSERT INTO hr_employees VALUES ('+','.join('?' for _ in cols)+')',[[r[c] for c in cols] for r in records])
    con.commit();return con

def run_queries(con,query_file,folder):
    results={}
    for block in query_file.read_text(encoding='utf-8').split('-- result: ')[1:]:
        name,sql=block.split('\n',1);sql=sql.strip()
        cur=con.execute(sql);heads=[d[0] for d in cur.description];rows=cur.fetchall()
        results[name.strip()]=[dict(zip(heads,r)) for r in rows]
        with (folder/(name.strip()+'.csv')).open('w',newline='',encoding='utf-8') as f:
            writer=csv.writer(f);writer.writerow(heads);writer.writerows(rows)
    return results

def verify(records,results):
    overview=results['overview'][0];n=len(records);exits=sum(r['attrition_flag'] for r in records)
    checks={'row_count':overview['employees']==n,'exit_count':overview['attrition_count']==exits,
            'active_reconciliation':overview['active_employees']+exits==n,
            'overall_denominator':abs(overview['attrition_rate_pct']-100*exits/n)<0.005,
            'unique_employee_ids':len({r['employee_id'] for r in records})==n}
    for group in ('department','job_role','overtime','age_band','tenure_band','job_satisfaction','department_overtime'):
        rows=results[group]
        checks[group+'_population']=sum(r['employees'] for r in rows)==n
        checks[group+'_exits']=sum(r['attrition_count'] for r in rows)==exits
        checks[group+'_rates']=all(abs(r['attrition_rate_pct']-100*r['attrition_count']/r['employees'])<0.005 for r in rows)
    for row in results['role_satisfaction']:
        expected=[r for r in records if r['job_role']==row['job_role'] and r['job_satisfaction']==row['job_satisfaction']]
        checks[f"cell_{row['job_role']}_{row['job_satisfaction']}"]=row['employees']==len(expected) and row['attrition_count']==sum(r['attrition_flag'] for r in expected)
    if not all(checks.values()):raise AssertionError([k for k,v in checks.items() if not v])
    return {'passed':sum(checks.values()),'failed':0,'checks':checks}

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,default=ROOT/'HR Data.xlsx')
    args=parser.parse_args();records,quality=prepare(args.input)
    (ROOT/'data').mkdir(exist_ok=True);(ROOT/'results').mkdir(exist_ok=True)
    with (ROOT/'data/hr_analytics.csv').open('w',newline='',encoding='utf-8') as f:
        wr=csv.DictWriter(f,fieldnames=list(records[0]));wr.writeheader();wr.writerows(records)
    con=load_sqlite(records,ROOT/'data/hr_analytics.sqlite')
    results=run_queries(con,ROOT/'sql/hr_analysis.sql',ROOT/'results')
    validation=verify(records,results);con.close()
    (ROOT/'results/validation.json').write_text(json.dumps({'data_quality':quality,**validation},indent=2),encoding='utf-8')
    (ROOT/'results/metrics.json').write_text(json.dumps(results,indent=2),encoding='utf-8')
    print(json.dumps({'overview':results['overview'][0],'checks_passed':validation['passed'],'checks_failed':0},indent=2))
if __name__=='__main__':main()
